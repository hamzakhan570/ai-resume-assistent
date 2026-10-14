"""ATS Resume Checker - Streamlit + Gemini Flash."""
import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_CHARS = 15000  # keep prompts small and fast

PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Analyze the resume below{jd_clause}. Be honest, specific and strict; do not inflate scores.

Return ONLY valid JSON with exactly this structure:
{{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting_and_structure": <integer 0-100>,
    "keywords_and_skills": <integer 0-100>,
    "work_experience_impact": <integer 0-100>,
    "education_and_certifications": <integer 0-100>,
    "readability_and_grammar": <integer 0-100>
  }},
  "strengths": ["<short point>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "issue": "<what is wrong>", "fix": "<exactly how to fix it>"}}
  ],
  "missing_keywords": ["<keyword>", "..."],
  "rewrite_suggestions": [
    {{"original": "<weak line from the resume>", "improved": "<stronger rewritten line>"}}
  ]
}}

Scoring guide: 90+ is exceptional, 75-89 is strong, 60-74 is average, below 60 needs major work.
Give 3-6 strengths, 5-10 improvements, up to 10 missing keywords and 2-4 rewrite suggestions.
{jd_block}
RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ---------- File reading ----------
def extract_text(uploaded_file) -> str:
    """Return plain text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ---------- AI analysis ----------
def parse_json(raw: str) -> dict:
    """Parse model output as JSON, tolerating code fences or stray text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    client = genai.Client(api_key=api_key)
    has_jd = bool(job_description.strip())
    prompt = PROMPT.format(
        jd_clause=" against the target job description" if has_jd else " for general ATS compatibility",
        jd_block=f'\nJOB DESCRIPTION:\n"""\n{job_description.strip()[:MAX_CHARS]}\n"""\n' if has_jd else "",
        resume=resume_text[:MAX_CHARS],
    )
    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    result = parse_json(response.text)

    # Normalise so the UI never crashes on missing/odd fields
    result["overall_score"] = clamp_score(result.get("overall_score"))
    result["section_scores"] = {
        k: clamp_score(v) for k, v in (result.get("section_scores") or {}).items()
    }
    for key in ("strengths", "improvements", "missing_keywords", "rewrite_suggestions"):
        if not isinstance(result.get(key), list):
            result[key] = []
    result["summary"] = str(result.get("summary", ""))
    return result


# ---------- UI helpers ----------
def score_color(score: int) -> str:
    if score >= 75:
        return "green"
    if score >= 60:
        return "orange"
    return "red"


def get_api_key() -> str:
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    return os.environ.get("GEMINI_API_KEY", "")


def render_results(result: dict) -> None:
    score = result["overall_score"]
    st.subheader("Your ATS Score")
    c1, c2 = st.columns([1, 3])
    with c1:
        st.markdown(f"## :{score_color(score)}[{score}/100]")
    with c2:
        st.progress(score / 100)
        st.write(result["summary"])

    if result["section_scores"]:
        st.subheader("Score breakdown")
        for name, val in result["section_scores"].items():
            st.write(f"**{name.replace('_', ' ').title()}** — {val}/100")
            st.progress(val / 100)

    if result["strengths"]:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.write(f"- {s}")

    if result["improvements"]:
        st.subheader("🛠️ Improvements")
        order = {"high": 0, "medium": 1, "low": 2}
        items = [i for i in result["improvements"] if isinstance(i, dict)]
        items.sort(key=lambda i: order.get(str(i.get("priority", "")).lower(), 3))
        icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
        for i in items:
            pr = str(i.get("priority", "")).lower()
            with st.expander(f"{icons.get(pr, '⚪')} {i.get('priority', '')} — {i.get('issue', '')}"):
                st.write(i.get("fix", ""))

    if result["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    rewrites = [r for r in result["rewrite_suggestions"] if isinstance(r, dict)]
    if rewrites:
        st.subheader("✍️ Suggested rewrites")
        for r in rewrites:
            st.markdown(f"**Before:** {r.get('original', '')}")
            st.markdown(f"**After:** {r.get('improved', '')}")
            st.divider()


# ---------- Main app ----------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.caption("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste a job description to check how well your resume matches it.",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
            if len(text) < 100:
                st.error("Couldn't read enough text. If your PDF is a scanned image, "
                         "upload a text-based PDF or a DOCX instead.")
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key, model.strip() or DEFAULT_MODEL)
            render_results(result)
        except Exception as e:  # show a friendly message instead of a stack trace
            st.error(f"Something went wrong: {e}")


if __name__ == "__main__":
    main()
